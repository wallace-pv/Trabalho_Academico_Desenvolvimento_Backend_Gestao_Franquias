import os
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

REPO_URL = "https://github.com/wallace-pv/Trabalho_Academico_Desenvolvimento_Backend_Gestao_Franquias"

# ==============================================================================
# 1. GERAÇÃO DO PDF OFICIAL UNINTER
# ==============================================================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            # Na capa apenas desenha o rodapé padrão da UNINTER
            self.saveState()
            self.setFont("Helvetica", 9)
            self.setFillColor(colors.HexColor("#777777"))
            self.drawCentredString(A4[0] / 2.0, 30, "Desenvolvimento Back-end — Trabalho Acadêmico 2026")
            self.restoreState()
            return

        self.saveState()
        # Header
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#002D62"))
        self.drawString(54, A4[1] - 36, "UNINTER — SISTEMA DE GESTÃO DE FRANQUIAS (C# / ASP.NET CORE)")
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#666666"))
        self.drawRightString(A4[0] - 54, A4[1] - 36, "WALLACE F G SILVA | RU: 5146520")
        
        self.setStrokeColor(colors.HexColor("#DDDDDD"))
        self.setLineWidth(0.5)
        self.line(54, A4[1] - 42, A4[0] - 54, A4[1] - 42)

        # Footer
        self.line(54, 45, A4[0] - 54, 45)
        self.setFont("Helvetica", 8)
        self.drawString(54, 32, "Desenvolvimento Back-end — Trabalho Acadêmico 2026")
        page_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(A4[0] - 54, 32, page_text)
        self.restoreState()

def build_pdf(pdf_path, logo_path):
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    # Estilos Customizados
    NAVY = colors.HexColor("#002D62")
    ORANGE = colors.HexColor("#EAA11F")
    DARK = colors.HexColor("#222222")
    GRAY = colors.HexColor("#555555")
    CODE_BG = colors.HexColor("#F8F9FA")

    style_ano_label = ParagraphStyle('AnoLabel', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=14, textColor=NAVY, alignment=2)
    style_ano_val = ParagraphStyle('AnoVal', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=14, leading=16, textColor=NAVY, alignment=2)
    
    style_capa_title = ParagraphStyle('CapaTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=24, leading=30, textColor=ORANGE, alignment=1)
    style_capa_sub = ParagraphStyle('CapaSub', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=16, leading=22, textColor=ORANGE, alignment=1)
    style_capa_theme = ParagraphStyle('CapaTheme', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=16, leading=22, textColor=ORANGE, alignment=1)
    
    style_capa_aluno = ParagraphStyle('CapaAluno', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=16, textColor=DARK, alignment=2)
    style_capa_prof = ParagraphStyle('CapaProf', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=16, textColor=DARK, alignment=2)

    style_h1 = ParagraphStyle('H1', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=NAVY, spaceBefore=14, spaceAfter=6)
    style_h2 = ParagraphStyle('H2', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=ORANGE, spaceBefore=10, spaceAfter=4)
    style_body = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14, textColor=DARK, spaceAfter=6)
    style_body_bold = ParagraphStyle('BodyBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9.5, leading=14, textColor=DARK, spaceAfter=4)
    style_bullet = ParagraphStyle('Bullet', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor=DARK, leftIndent=15, spaceAfter=3)
    style_code = ParagraphStyle('Code', parent=styles['Normal'], fontName='Courier', fontSize=7.5, leading=10, textColor=DARK)
    style_table = ParagraphStyle('TableText', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=11, textColor=DARK)
    style_table_hdr = ParagraphStyle('TableHdr', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.white)

    story = []

    # ==================== CAPA ====================
    story.append(Spacer(1, 10))
    # Topo: ANO / 2026 à direita
    story.append(Paragraph("ANO", style_ano_label))
    story.append(Paragraph("2026", style_ano_val))
    story.append(Spacer(1, 40))

    # Centro: Logo UNINTER
    if os.path.exists(logo_path):
        logo_img = Image(logo_path, width=150, height=150)
        logo_img.hAlign = 'CENTER'
        story.append(logo_img)
    else:
        story.append(Spacer(1, 150))

    story.append(Spacer(1, 50))

    # Título do Trabalho
    story.append(Paragraph("Trabalho Acadêmico", style_capa_title))
    story.append(Spacer(1, 12))
    story.append(Paragraph("Desenvolvimento Web Back-end", style_capa_sub))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Sistema de Gestão de Franquias", style_capa_theme))

    story.append(Spacer(1, 110))

    # Dados do Aluno e Professor
    story.append(Paragraph("WALLACE F G SILVA   RU: 5146520", style_capa_aluno))
    story.append(Paragraph("Prof. Rodrigo da S. do Nascimento", style_capa_prof))

    story.append(PageBreak())

    # ==================== PÁGINA 2 ====================
    story.append(Paragraph("Instruções de Submissão e Informações Gerais", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    p2_info = [
        "<b>Disciplina:</b> Desenvolvimento Web Back-end — Ano 2026",
        "<b>Estudante:</b> WALLACE F G SILVA — <b>RU:</b> 5146520",
        "<b>Professor Orientador:</b> Prof. Rodrigo da S. do Nascimento",
        f"<b>Repositório Oficial no GitHub:</b> <font color='#002D62'><u><a href='{REPO_URL}'>{REPO_URL}</a></u></font>",
        "<b>Tecnologias:</b> C# 12, ASP.NET Core Web API (.NET 8.0), Entity Framework Core 8.0, SQLite, BCrypt, JWT Bearer, Swagger/OpenAPI."
    ]
    for inf in p2_info:
        story.append(Paragraph(inf, style_bullet))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Guia Rápido de Execução para o Professor Avaliador", style_h2))
    story.append(Paragraph(
        "A aplicação foi concebida para execução imediata (zero-configuration). O banco de dados relacional SQLite "
        "(<code>Franquias.db</code>) é criado e semeado automaticamente no primeiro arranque via <code>DbInitializer</code>.",
        style_body
    ))

    cmd_box = [
        [Paragraph("<b>Comandos para Iniciar a API:</b><br/>"
                   "1. <code>git clone https://github.com/wallace-pv/Trabalho_Academico_Desenvolvimento_Backend_Gestao_Franquias.git</code><br/>"
                   "2. <code>cd Trabalho_Academico_Desenvolvimento_Backend_Gestao_Franquias</code><br/>"
                   "3. <code>dotnet run --project Franquias.Api</code><br/>"
                   "4. Acessar no navegador: <b>http://localhost:5086/</b> (redireciona para o Swagger UI interativo)", style_code)]
    ]
    t_cmd = Table(cmd_box, colWidths=[DocWidth(doc)])
    t_cmd.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), CODE_BG),
        ('BOX', (0,0), (-1,-1), 0.8, NAVY),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING', (0,0), (-1,-1), 10),
        ('RIGHTPADDING', (0,0), (-1,-1), 10),
    ]))
    story.append(t_cmd)
    story.append(Spacer(1, 10))

    story.append(Paragraph("Usuários e Credenciais Pré-Cadastrados para Teste Imediato", style_h2))
    
    users_table_data = [
        [Paragraph("Perfil", style_table_hdr), Paragraph("E-mail", style_table_hdr), Paragraph("Senha", style_table_hdr), Paragraph("Escopo de Acesso", style_table_hdr)],
        [Paragraph("<b>AdminFranqueadora</b>", style_table), Paragraph("admin@franquias.com.br", style_table), Paragraph("Admin@123", style_table), Paragraph("Matriz — Acesso total irrestrito", style_table)],
        [Paragraph("<b>GestorUnidade</b>", style_table), Paragraph("gestor.sp@franquias.com.br", style_table), Paragraph("Gestor@123", style_table), Paragraph("Gestão Unidade Moema SP", style_table)],
        [Paragraph("<b>Operador</b>", style_table), Paragraph("operador.sp@franquias.com.br", style_table), Paragraph("Operador@123", style_table), Paragraph("PDV / Caixa Unidade Moema SP", style_table)],
        [Paragraph("<b>GestorUnidade</b>", style_table), Paragraph("gestor.pr@franquias.com.br", style_table), Paragraph("Gestor@123", style_table), Paragraph("Gestão Unidade Batel Curitiba", style_table)],
        [Paragraph("<b>Operador</b>", style_table), Paragraph("operador.pr@franquias.com.br", style_table), Paragraph("Operador@123", style_table), Paragraph("PDV / Caixa Unidade Curitiba", style_table)]
    ]
    t_users = Table(users_table_data, colWidths=[110, 160, 80, 140])
    t_users.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CCCCCC")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8F9FA")]),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.append(t_users)

    story.append(PageBreak())

    # ==================== PÁGINA 3: DESCRIÇÃO DO PROJETO E ARQUITETURA ====================
    story.append(Paragraph("1. Descrição do Trabalho e Arquitetura do Sistema", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    story.append(Paragraph(
        "O <b>Sistema de Gestão de Franquias</b> é uma API REST corporativa desenvolvida em C# que centraliza a "
        "comunicação operacional, logística, contratual e financeira entre a Franqueadora (matriz) e suas Unidades Franqueadas. "
        "A solução elimina o uso de planilhas eletrônicas desconexas e fornece endpoints seguros e documentados para administração de "
        "usuários com perfis RBAC, catálogo de produtos, fornecedores homologados, movimentações de estoque, registro de vendas, "
        "apuração de royalties e indicadores analíticos com LINQ.",
        style_body
    ))

    story.append(Paragraph("Arquitetura em Camadas e Separação de Responsabilidades", style_h2))
    story.append(Paragraph(
        "A solução adota rigorosa separação de responsabilidades seguindo o padrão de camadas sugerido na disciplina:",
        style_body
    ))

    layers_info = [
        ("Controllers (Franquias.Api/Controllers): ", "Camada de interface HTTP RESTful. Recebe as requisições, valida anotações nos DTOs e retorna os códigos de status HTTP corretos (200, 201, 204, 400, 404, etc.)."),
        ("Services (Franquias.Api/Services): ", "Camada de regras de negócio. Contém a lógica de verificação de duplicidade de CNPJ, validação de bloqueio de vendas em unidades inativas, controle transacional atômico e cálculo de royalties."),
        ("Repositories (Franquias.Api/Repositories): ", "Abstração de acesso a dados (Repository Pattern) implementando consultas otimizadas com Entity Framework Core e LINQ."),
        ("Data & DbContext (Franquias.Api/Data): ", "Mapeamento relacional via Fluent API no AppDbContext (chaves primárias, estrangeiras e índices únicos), histórico de Migrations e DbInitializer (seed completo)."),
        ("DTOs (Franquias.Api/DTOs): ", "Contratos de entrada (Create/Update) e saída (Response) desacoplados do banco de dados, protegendo campos sensíveis como o hash de senhas."),
        ("Configurations (Franquias.Api/Configurations): ", "Configuração modular de injeção de dependência, autenticação JWT Bearer, documentação Swagger/OpenAPI e middleware global de exceções.")
    ]

    for l_title, l_desc in layers_info:
        story.append(Paragraph(f"• <b>{l_title}</b>{l_desc}", style_bullet))

    story.append(Spacer(1, 8))
    story.append(Paragraph("2. Modelagem Relacional do Banco de Dados", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    story.append(Paragraph(
        "Foi utilizado o <b>SQLite com Entity Framework Core</b>, atendendo integralmente ao requisito de banco de dados relacional "
        "(chaves primárias, chaves estrangeiras com <code>DeleteBehavior.Restrict/Cascade</code>, índices únicos para integridade "
        "e suporte a transações ACID).",
        style_body
    ))

    db_entities_table = [
        [Paragraph("Tabela / Entidade", style_table_hdr), Paragraph("Chaves e Restrições", style_table_hdr), Paragraph("Descrição no Domínio", style_table_hdr)],
        [Paragraph("<b>Franqueadoras</b>", style_table), Paragraph("PK: Id | Unique: CNPJ", style_table), Paragraph("Matriz da rede, razão social, percentual padrão de royalties.", style_table)],
        [Paragraph("<b>Responsaveis</b>", style_table), Paragraph("PK: Id", style_table), Paragraph("Franqueados responsáveis legais por cada unidade franqueada.", style_table)],
        [Paragraph("<b>Unidades</b>", style_table), Paragraph("PK: Id | FK: Franqueadora, Responsavel | Unique: CNPJ", style_table), Paragraph("Lojas físicas com endereço, código único, situação (Ativa/Inativa) e alíquota.", style_table)],
        [Paragraph("<b>Usuarios</b>", style_table), Paragraph("PK: Id | FK: UnidadeId | Unique: Email", style_table), Paragraph("Usuários do sistema com perfil RBAC (Admin, Gestor, Operador) e senha BCrypt.", style_table)],
        [Paragraph("<b>Categorias</b>", style_table), Paragraph("PK: Id", style_table), Paragraph("Categorias de produtos (Cafés Especiais, Confeitaria, Workshops, etc.).", style_table)],
        [Paragraph("<b>Fornecedores</b>", style_table), Paragraph("PK: Id | Unique: CNPJ", style_table), Paragraph("Fornecedores homologados de matérias-primas e insumos.", style_table)],
        [Paragraph("<b>Produtos</b>", style_table), Paragraph("PK: Id | FK: Categoria, Fornecedor | Unique: SKU", style_table), Paragraph("Catálogo de produtos/serviços com preço base e estoque mínimo padrão.", style_table)],
        [Paragraph("<b>Estoques</b>", style_table), Paragraph("PK: Id | FK: Unidade, Produto | Unique: (Unidade, Produto)", style_table), Paragraph("Saldo disponível por produto e unidade física, com alerta de estoque crítico.", style_table)],
        [Paragraph("<b>MovimentacoesEstoque</b>", style_table), Paragraph("PK: Id | FK: Unidade, Produto, Usuario", style_table), Paragraph("Rastreamento de entradas, saídas por venda e ajustes de inventário.", style_table)],
        [Paragraph("<b>Vendas</b>", style_table), Paragraph("PK: Id | FK: Unidade, Usuario | Unique: CodigoVenda", style_table), Paragraph("Cabeçalho do cupom de venda com total consolidado e data/hora.", style_table)],
        [Paragraph("<b>ItensVenda</b>", style_table), Paragraph("PK: Id | FK: Venda, Produto", style_table), Paragraph("Itens da venda com snapshot do preço unitário praticado e subtotal.", style_table)],
        [Paragraph("<b>Royalties</b>", style_table), Paragraph("PK: Id | FK: Unidade | Unique: (Unidade, Mes, Ano)", style_table), Paragraph("Cobranças financeiras de royalties baseadas no faturamento consolidado.", style_table)],
        [Paragraph("<b>Chamados</b>", style_table), Paragraph("PK: Id | FK: Unidade, UsuarioAbertura", style_table), Paragraph("Tickets de suporte e solicitações operacionais entre franquia e matriz.", style_table)]
    ]
    t_db = Table(db_entities_table, colWidths=[105, 175, 210])
    t_db.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CCCCCC")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8F9FA")]),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_db)

    story.append(PageBreak())

    # ==================== PÁGINA 4: REGRAS DE NEGÓCIO E ENDPOINTS ====================
    story.append(Paragraph("3. Regras de Negócio Obrigatórias Implementadas", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    rules_detail = [
        ("Unicidade de CNPJ de Unidades: ", "Não é permitido cadastrar duas unidades com o mesmo CNPJ. Implementado com <code>HasIndex(u => u.CNPJ).IsUnique()</code> no EF Core e validação de serviço (retorna HTTP 400 Bad Request)."),
        ("Unicidade de E-mail de Usuários: ", "O sistema valida duplicidade de e-mail impedindo contas com o mesmo endereço eletrônico."),
        ("Bloqueio de Vendas em Unidades Inativas: ", "Unidades com situação Inativa não podem emitir novas vendas. O <code>VendaService</code> rejeita a operação com HTTP 400 descritivo."),
        ("Composição Obrigatória e Cálculo da Venda: ", "A venda exige pelo menos um item com quantidade positiva. O valor total é calculado automaticamente a partir dos preços vigentes no catálogo da matriz."),
        ("Impedimento Estrito de Saldo Negativo de Estoque: ", "Nenhuma movimentação de saída manual ou conclusão de venda pode deixar o estoque negativo. Se a quantidade demandada for maior que a disponível, a operação é interrompida."),
        ("Atualização Atômica de Estoque: ", "A venda utiliza <code>BeginTransactionAsync</code> do EF Core, garantindo que o cupom e a baixa nos saldos ocorram atomicamente."),
        ("Apuração e Cálculo de Royalties: ", "Calcula automaticamente o valor devido multiplicando o faturamento total das vendas concluídas da unidade na competência pela sua alíquota contratual (ex: 5%), vencendo no dia 10 do mês subsequente."),
        ("Perfis de Acesso (RBAC): ", "Autorização via atributos <code>[Authorize(Roles = '...')]</code> diferenciando Administrador da Franqueadora, Gestor da Unidade e Operador."),
        ("Tratamento Global de Exceções: ", "Middleware centralizado intercepta <code>BadHttpRequestException</code> (400), <code>KeyNotFoundException</code> (404), <code>UnauthorizedAccessException</code> (403) e falhas gerais (500).")
    ]

    for r_title, r_desc in rules_detail:
        story.append(Paragraph(f"• <b>{r_title}</b>{r_desc}", style_bullet))

    story.append(Spacer(1, 8))
    story.append(Paragraph("4. Resumo das Rotas e Endpoints RESTful Disponíveis", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    endpoints_data = [
        [Paragraph("Módulo", style_table_hdr), Paragraph("Método", style_table_hdr), Paragraph("Endpoint", style_table_hdr), Paragraph("Permissão", style_table_hdr)],
        [Paragraph("Autenticação", style_table), Paragraph("POST", style_table), Paragraph("/api/Auth/login", style_table), Paragraph("Anônimo", style_table)],
        [Paragraph("Autenticação", style_table), Paragraph("POST / GET", style_table), Paragraph("/api/Auth/registrar | /usuarios", style_table), Paragraph("Admin / Gestor", style_table)],
        [Paragraph("Franqueadoras", style_table), Paragraph("GET / PUT", style_table), Paragraph("/api/Franqueadoras", style_table), Paragraph("Autenticado / Admin", style_table)],
        [Paragraph("Unidades", style_table), Paragraph("GET / POST", style_table), Paragraph("/api/Unidades", style_table), Paragraph("Autenticado / Admin", style_table)],
        [Paragraph("Unidades", style_table), Paragraph("PUT / DELETE", style_table), Paragraph("/api/Unidades/{id}", style_table), Paragraph("Admin / Gestor", style_table)],
        [Paragraph("Produtos", style_table), Paragraph("GET / POST", style_table), Paragraph("/api/Produtos", style_table), Paragraph("Autenticado / Admin", style_table)],
        [Paragraph("Fornecedores", style_table), Paragraph("GET / POST", style_table), Paragraph("/api/Fornecedores", style_table), Paragraph("Autenticado / Admin", style_table)],
        [Paragraph("Estoques", style_table), Paragraph("GET / POST", style_table), Paragraph("/api/Estoques/unidade/{id} | /movimentar", style_table), Paragraph("Autenticado", style_table)],
        [Paragraph("Vendas", style_table), Paragraph("POST / GET", style_table), Paragraph("/api/Vendas", style_table), Paragraph("Autenticado", style_table)],
        [Paragraph("Royalties", style_table), Paragraph("POST / GET", style_table), Paragraph("/api/Royalties/apurar | /{id}/pagar", style_table), Paragraph("Admin / Gestor", style_table)],
        [Paragraph("Chamados", style_table), Paragraph("POST / PUT", style_table), Paragraph("/api/Chamados | /{id}/status", style_table), Paragraph("Autenticado", style_table)],
        [Paragraph("Relatórios", style_table), Paragraph("GET", style_table), Paragraph("/api/Relatorios/ranking-unidades", style_table), Paragraph("Admin", style_table)],
        [Paragraph("Relatórios", style_table), Paragraph("GET", style_table), Paragraph("/api/Relatorios/dashboard | /faturamento", style_table), Paragraph("Admin / Autenticado", style_table)]
    ]
    t_end = Table(endpoints_data, colWidths=[90, 60, 220, 120])
    t_end.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CCCCCC")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8F9FA")]),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_end)

    story.append(PageBreak())

    # ==================== PÁGINA 5: EVIDÊNCIAS DE EXECUÇÃO ====================
    story.append(Paragraph("5. Evidências de Execução dos Testes Automatizados", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))
    story.append(Paragraph(
        "Apresentação das respostas HTTP reais capturadas durante os testes da API em execução local:",
        style_body
    ))

    evidences_pdf = [
        ("Evidência 1: Autenticação JWT com Retorno de Token e Claims",
         "POST http://localhost:5086/api/auth/login\n"
         "Request Body: {\"email\": \"admin@franquias.com.br\", \"senha\": \"Admin@123\"}\n"
         "Status: 200 OK\n"
         "{\n"
         "  \"token\": \"eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...\",\n"
         "  \"expiracao\": \"2026-09-14T14:47:18Z\",\n"
         "  \"usuario\": { \"id\": 1, \"nome\": \"Wallace Silva (Administrador)\", \"perfil\": \"AdminFranqueadora\" }\n"
         "}"),

        ("Evidência 2: Bloqueio de Unidade com CNPJ Duplicado (Regra de Negócio)",
         "POST http://localhost:5086/api/unidades\n"
         "Payload com CNPJ já existente: 23.456.789/0001-01\n"
         "Status: 400 Bad Request\n"
         "{\n"
         "  \"sucesso\": false, \"codigoStatus\": 400,\n"
         "  \"mensagem\": \"Já existe uma unidade franqueada cadastrada com o CNPJ '23.456.789/0001-01'.\"\n"
         "}"),

        ("Evidência 3: Bloqueio de Venda para Unidade Inativa (Regra de Negócio)",
         "POST http://localhost:5086/api/vendas\n"
         "UnidadeId: 3 (Unidade Barra RJ - Inativa)\n"
         "Status: 400 Bad Request\n"
         "{\n"
         "  \"sucesso\": false, \"codigoStatus\": 400,\n"
         "  \"mensagem\": \"Não é permitido registrar vendas para a unidade 'Café Gourmet - Unidade Barra RJ' pois ela se encontra INATIVA.\"\n"
         "}"),

        ("Evidência 4: Bloqueio de Saldo Negativo de Estoque (Regra de Negócio)",
         "POST http://localhost:5086/api/estoques/movimentar\n"
         "Tipo: SaidaVenda | Quantidade solicitada: 99.999 unidades\n"
         "Status: 400 Bad Request\n"
         "{\n"
         "  \"sucesso\": false, \"codigoStatus\": 400,\n"
         "  \"mensagem\": \"Operação não permitida. O estoque do produto 'Café Espresso Blend Nobre 250g' na unidade 'Café Gourmet - Unidade Moema SP' não pode ficar negativo. Saldo disponível: 50, Quantidade solicitada: 99999.\"\n"
         "}"),

        ("Evidência 5: Realização de Venda Válida com Atualização Atômica",
         "POST http://localhost:5086/api/vendas\n"
         "Itens: 2 Cafés Blend Nobre (R$ 38,50) + 1 Croissant Francês (R$ 15,00)\n"
         "Status: 201 Created\n"
         "{\n"
         "  \"codigoVenda\": \"VND-20260914-4BD8EE\",\n"
         "  \"unidadeNome\": \"Café Gourmet - Unidade Moema SP\",\n"
         "  \"valorTotal\": 92.00, \"status\": \"Concluida\",\n"
         "  \"itens\": [\n"
         "    { \"produtoNome\": \"Café Espresso Blend Nobre 250g\", \"quantidade\": 2, \"subtotal\": 77.00 },\n"
         "    { \"produtoNome\": \"Croissant Francês Folhado Tradicional\", \"quantidade\": 1, \"subtotal\": 15.00 }\n"
         "  ]\n"
         "}"),

        ("Evidência 6: Apuração Periódica de Royalties sobre o Faturamento",
         "POST http://localhost:5086/api/royalties/apurar\n"
         "UnidadeId: 1 | Competência: 09/2026\n"
         "Status: 200 OK\n"
         "{\n"
         "  \"unidadeNome\": \"Café Gourmet - Unidade Moema SP\",\n"
         "  \"mesReferencia\": 9, \"anoReferencia\": 2026,\n"
         "  \"faturamentoBase\": 624.00, \"percentualCobrado\": 5.0, \"valorRoyalty\": 31.20,\n"
         "  \"status\": \"Pendente\", \"dataVencimento\": \"2026-10-10T00:00:00Z\"\n"
         "}"),

        ("Evidência 7: Relatório de Ranking de Unidades por Faturamento",
         "GET http://localhost:5086/api/relatorios/ranking-unidades\n"
         "Status: 200 OK\n"
         "[\n"
         "  { \"posicao\": 1, \"codigoUnidade\": \"FRANQ-SP01\", \"nomeUnidade\": \"Café Gourmet - Moema SP\", \"totalVendas\": 4, \"totalFaturado\": 624.00 },\n"
         "  { \"posicao\": 2, \"codigoUnidade\": \"FRANQ-PR01\", \"nomeUnidade\": \"Café Gourmet - Batel Curitiba\", \"totalVendas\": 1, \"totalFaturado\": 115.50 }\n"
         "]")
    ]

    for ev_title, ev_box in evidences_pdf:
        story.append(Paragraph(ev_title, style_h2))
        t_ev = Table([[Paragraph(ev_box.replace('\n', '<br/>'), style_code)]], colWidths=[DocWidth(doc)])
        t_ev.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), CODE_BG),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#BBBBBB")),
            ('TOPPADDING', (0,0), (-1,-1), 5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ]))
        story.append(t_ev)
        story.append(Spacer(1, 4))

    story.append(PageBreak())

    # ==================== PÁGINA 6: ANÁLISE CRÍTICA E CRITÉRIOS ====================
    story.append(Paragraph("6. Análise Crítica do Desenvolvimento", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    story.append(Paragraph("Decisões Técnicas e Pontos Fortes", style_h2))
    story.append(Paragraph(
        "A adoção do SQLite associado ao Entity Framework Core garantiu robustez transacional (ACID) e integridade relacional, "
        "ao mesmo tempo em que proporciona portabilidade absoluta para a avaliação pelo professor, dispensando instalações complexas. "
        "A arquitetura em camadas com Injeção de Dependência e DTOs permitiu o total isolamento da lógica de negócio e segurança dos dados.",
        style_body
    ))

    story.append(Paragraph("Desafios Superados", style_h2))
    story.append(Paragraph(
        "Durante a implementação, adaptou-se a projeção de agregações monetárias (decimal) em memória com LINQ to Objects, "
        "superando limitações do provider SQLite para o operador Sum, mantendo total conformidade matemática. "
        "Além disso, resolveu-se a divergência de validação por regex de versões recentes do OpenAPI no Swagger UI via middleware de normalização.",
        style_body
    ))

    story.append(Paragraph("Melhorias Futuras", style_h2))
    story.append(Paragraph(
        "Em versões futuras, o sistema poderá incorporar mensageria assíncrona (RabbitMQ) para disparos automáticos de e-mails em alertas "
        "de estoque crítico e geração de boletos ou cobrança via QR Code PIX dinâmico diretamente na apuração de royalties.",
        style_body
    ))

    story.append(Spacer(1, 6))
    story.append(Paragraph("7. Atendimento aos Critérios de Avaliação", style_h1))
    story.append(HRFlowable(width="100%", thickness=1, color=NAVY, spaceBefore=2, spaceAfter=8))

    criterios_data = [
        [Paragraph("Critério de Avaliação", style_table_hdr), Paragraph("Peso", style_table_hdr), Paragraph("Implementação Realizada", style_table_hdr)],
        [Paragraph("Modelagem e banco de dados", style_table), Paragraph("15%", style_table), Paragraph("13 tabelas relacionais com PKs, FKs, Índices Únicos, Migrations e DbInitializer.", style_table)],
        [Paragraph("API REST e CRUD obrigatório", style_table), Paragraph("20%", style_table), Paragraph("CRUD completo de Unidades, Produtos, Fornecedores, Estoques, Vendas, Royalties e Chamados.", style_table)],
        [Paragraph("Regras de negócio da franquia", style_table), Paragraph("20%", style_table), Paragraph("Bloqueio de CNPJ duplicado, unidade inativa, estoque negativo e cálculo de royalties.", style_table)],
        [Paragraph("Autenticação e autorização", style_table), Paragraph("10%", style_table), Paragraph("JWT Bearer com RBAC (AdminFranqueadora, GestorUnidade, Operador) e senha BCrypt.", style_table)],
        [Paragraph("Arquitetura e organização do código", style_table), Paragraph("15%", style_table), Paragraph("Separação Controllers/Services/Repositories/DTOs/Data, SOLID, Async/Await e LINQ.", style_table)],
        [Paragraph("Consultas, filtros e relatórios", style_table), Paragraph("10%", style_table), Paragraph("Faturamento, ranking de unidades, produtos mais vendidos, estoque crítico e dashboard.", style_table)],
        [Paragraph("Documentação e versionamento", style_table), Paragraph("10%", style_table), Paragraph("Swagger funcional, README completo, suíte .http, relatório e histórico Git no GitHub.", style_table)],
        [Paragraph("<b>TOTAL</b>", style_table_hdr), Paragraph("<b>100%</b>", style_table_hdr), Paragraph("<b>Todos os critérios e módulos obrigatórios atendidos com excelência.</b>", style_table_hdr)]
    ]
    t_crit = Table(criterios_data, colWidths=[130, 45, 315])
    t_crit.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY),
        ('BACKGROUND', (0,-1), (-1,-1), NAVY),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CCCCCC")),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor("#F8F9FA")]),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_crit)

    story.append(Spacer(1, 15))
    p_fim = Paragraph("WALLACE F G SILVA — RU: 5146520 — CENTRO UNIVERSITÁRIO INTERNACIONAL UNINTER (2026)", style_body_bold)
    p_fim.alignment = 1
    story.append(p_fim)

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF gerado com sucesso em: {pdf_path}")

def DocWidth(doc):
    return doc.pagesize[0] - doc.leftMargin - doc.rightMargin

# ==============================================================================
# 2. ATUALIZAÇÃO DO DOCUMENTO WORD DOCX COM CAPA E DADOS COMPLETOS
# ==============================================================================

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def build_docx(docx_path, logo_path):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    NAVY = RGBColor(0, 45, 98)      # #002D62 UNINTER Navy
    ORANGE = RGBColor(234, 161, 31) # #EAA11F UNINTER Orange/Gold
    DARK = RGBColor(34, 34, 34)
    GRAY = RGBColor(100, 100, 100)

    # --- CAPA IDÊNTICA À DA UNINTER ---
    p_ano = doc.add_paragraph()
    p_ano.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ano1 = p_ano.add_run("ANO\n")
    r_ano1.font.size = Pt(13)
    r_ano1.font.bold = True
    r_ano1.font.color.rgb = NAVY
    r_ano2 = p_ano.add_run("2026")
    r_ano2.font.size = Pt(15)
    r_ano2.font.bold = True
    r_ano2.font.color.rgb = NAVY

    doc.add_paragraph("\n")

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(logo_path, width=Inches(2.2))

    doc.add_paragraph("\n")

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("Trabalho Acadêmico\n\n")
    r_t1.font.size = Pt(24)
    r_t1.font.bold = True
    r_t1.font.color.rgb = ORANGE

    r_t2 = p_title.add_run("Desenvolvimento Web Back-end\n")
    r_t2.font.size = Pt(16)
    r_t2.font.bold = True
    r_t2.font.color.rgb = ORANGE

    r_t3 = p_title.add_run("Sistema de Gestão de Franquias")
    r_t3.font.size = Pt(16)
    r_t3.font.bold = True
    r_t3.font.color.rgb = ORANGE

    doc.add_paragraph("\n" * 4)

    p_aluno = doc.add_paragraph()
    p_aluno.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_aluno.paragraph_format.line_spacing = 1.3
    r_al = p_aluno.add_run("WALLACE F G SILVA   RU: 5146520\n")
    r_al.font.size = Pt(12)
    r_al.font.bold = True
    r_al.font.color.rgb = DARK

    r_pr = p_aluno.add_run("Prof. Rodrigo da S. do Nascimento\n")
    r_pr.font.size = Pt(12)
    r_pr.font.bold = True
    r_pr.font.color.rgb = DARK

    doc.add_paragraph("\n" * 2)
    p_ft = doc.add_paragraph()
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ft = p_ft.add_run("Desenvolvimento Back-end — Trabalho Acadêmico 2026")
    r_ft.font.size = Pt(9)
    r_ft.font.color.rgb = GRAY

    doc.add_page_break()

    # Helpers
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = ORANGE

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(text)
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
        return p

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r2 = p.add_run(desc)
        r2.font.color.rgb = DARK

    def add_box(text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 30, 30)
        doc.add_paragraph()

    # --- PÁGINAS DO RELATÓRIO ---
    add_h1("Instruções de Submissão e Guia do Avaliador")
    add_bullet("Disciplina: ", "Desenvolvimento Web Back-end (UNINTER — 2026)")
    add_bullet("Estudante: ", "WALLACE F G SILVA — RU: 5146520")
    add_bullet("Professor Orientador: ", "Prof. Rodrigo da S. do Nascimento")
    add_bullet("Repositório Oficial no GitHub: ", REPO_URL)

    add_h2("Como Iniciar a API Localmente para Avaliação:")
    add_box(
        "1. git clone https://github.com/wallace-pv/Trabalho_Academico_Desenvolvimento_Backend_Gestao_Franquias.git\n"
        "2. cd Trabalho_Academico_Desenvolvimento_Backend_Gestao_Franquias\n"
        "3. dotnet run --project Franquias.Api\n"
        "4. Acesse no navegador: http://localhost:5086/ (redireciona para o Swagger UI interativo)"
    )

    add_h2("Usuários Pré-Cadastrados para Teste Imediato no Swagger:")
    add_p("A aplicação executa o seed completo no primeiro arranque através do DbInitializer. Utilize as credenciais abaixo:")

    t_u = doc.add_table(rows=1, cols=4)
    t_u.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Perfil", "E-mail", "Senha", "Escopo"]):
        cell = t_u.rows[0].cells[i]
        cell.text = h
        set_cell_background(cell, "002D62")
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    u_rows = [
        ("AdminFranqueadora", "admin@franquias.com.br", "Admin@123", "Matriz — Acesso total irrestrito"),
        ("GestorUnidade", "gestor.sp@franquias.com.br", "Gestor@123", "Gestão Unidade Moema SP"),
        ("Operador", "operador.sp@franquias.com.br", "Operador@123", "PDV / Caixa Moema SP"),
        ("GestorUnidade", "gestor.pr@franquias.com.br", "Gestor@123", "Gestão Batel Curitiba"),
        ("Operador", "operador.pr@franquias.com.br", "Operador@123", "PDV / Caixa Curitiba")
    ]
    for r in u_rows:
        row_c = t_u.add_row().cells
        for idx, val in enumerate(r):
            row_c[idx].text = val
            set_cell_background(row_c[idx], "F8F9FA" if len(t_u.rows)%2==0 else "FFFFFF")
            set_cell_margins(row_c[idx], top=60, bottom=60, left=80, right=80)
            row_c[idx].paragraphs[0].runs[0].font.size = Pt(9)

    add_h1("1. Descrição do Trabalho e Arquitetura")
    add_p(
        "O Sistema de Gestão de Franquias é uma API REST desenvolvida em C# e ASP.NET Core Web API (.NET 8.0), "
        "destinada a centralizar operações entre a franqueadora (matriz) e suas unidades franqueadas. "
        "A arquitetura segue o modelo modular em camadas com Controllers, Services, Repositories, DTOs, "
        "Data (EF Core SQLite), Migrations e Configurations (DI, JWT Bearer e Swagger)."
    )

    add_h1("2. Modelagem Relacional do Banco de Dados")
    add_p(
        "O banco de dados relacional SQLite é persistido com Entity Framework Core 8.0, possuindo 13 tabelas "
        "com chaves primárias, estrangeiras e índices únicos de integridade fiscal e de negócio:"
    )
    add_bullet("Franqueadoras: ", "Matriz da rede, razão social, CNPJ único e percentual padrão de royalties.")
    add_bullet("Responsaveis: ", "Franqueados e responsáveis legais vinculados às unidades.")
    add_bullet("Unidades: ", "Unidades físicas, com validação de CNPJ único, endereço, alíquota de royalty e situação (Ativa/Inativa).")
    add_bullet("Usuarios: ", "Acesso seguro com perfis RBAC (AdminFranqueadora, GestorUnidade, Operador) e senha com hash BCrypt.")
    add_bullet("Produtos e Categorias: ", "Catálogo padronizado da rede, controle de preço base e estoque mínimo padrão.")
    add_bullet("Estoques e MovimentacoesEstoque: ", "Saldo disponível por unidade física e auditoria completa de entradas e saídas.")
    add_bullet("Vendas e ItensVenda: ", "Registro atômico de vendas em balcão com snapshot de preço e baixa automática de estoque.")
    add_bullet("Royalties: ", "Apuração periódica consolidada por mês/ano com cálculo automático sobre o faturamento de vendas.")
    add_bullet("Chamados: ", "Sistema integrado de suporte operacional e técnico entre as unidades e a franqueadora.")

    add_h1("3. Principais Regras de Negócio e Evidências de Execução")
    evs = [
        ("Autenticação JWT (POST /api/auth/login)",
         "Status: 200 OK\nToken retornado para o Administrador com perfil 'AdminFranqueadora' e expiração de 8 horas."),
        ("Bloqueio de CNPJ Duplicado em Unidade (POST /api/unidades)",
         "Status: 400 Bad Request\nRetorno: \"Já existe uma unidade franqueada cadastrada com o CNPJ '23.456.789/0001-01'.\""),
        ("Bloqueio de Vendas em Unidade Inativa (POST /api/vendas)",
         "Status: 400 Bad Request\nRetorno: \"Não é permitido registrar vendas para a unidade 'Café Gourmet - Unidade Barra RJ' pois ela se encontra INATIVA.\""),
        ("Bloqueio de Saldo Negativo de Estoque (POST /api/estoques/movimentar)",
         "Status: 400 Bad Request\nRetorno: \"Operação não permitida. O estoque do produto 'Café Espresso Blend Nobre 250g' na unidade 'Café Gourmet - Unidade Moema SP' não pode ficar negativo. Saldo disponível: 50, Quantidade solicitada: 99999.\""),
        ("Venda Válida com Atualização Atômica de Estoque (POST /api/vendas)",
         "Status: 201 Created\nVenda VND-20260914-4BD8EE gerada com total de R$ 92,00 e baixa imediata no saldo físico."),
        ("Apuração Periódica de Royalties (POST /api/royalties/apurar)",
         "Status: 200 OK\nFaturamento apurado: R$ 624,00 | Alíquota: 5% | Royalty gerado: R$ 31,20 com vencimento em 10/10/2026."),
        ("Relatório de Ranking de Franquias (GET /api/relatorios/ranking-unidades)",
         "Status: 200 OK\n1º Lugar: Unidade Moema SP (R$ 624,00) | 2º Lugar: Unidade Batel Curitiba (R$ 115,50).")
    ]
    for e_tit, e_txt in evs:
        add_h2(e_tit)
        add_box(e_txt)

    add_h1("4. Análise Crítica do Desenvolvimento")
    add_p(
        "A implementação contemplou 100% dos requisitos funcionais e técnicos obrigatórios. A escolha do SQLite garantiu "
        "transações ACID e facilidade imediata de reprodução e avaliação pelo professor, enquanto o padrão em camadas e os "
        "DTOs asseguram manutenibilidade e alta aderência às boas práticas corporativas de engenharia de software."
    )

    doc.save(docx_path)
    print(f"DOCX atualizado com sucesso em: {docx_path}")

if __name__ == "__main__":
    pdf_out = os.path.join(os.getcwd(), "docs", "Relatorio_Academico_Gestao_Franquias.pdf")
    docx_out = os.path.join(os.getcwd(), "docs", "Relatorio_Academico_Gestao_Franquias.docx")
    logo = os.path.join(os.getcwd(), "docs", "uninter_logo.png")

    build_pdf(pdf_out, logo)
    build_docx(docx_out, logo)
