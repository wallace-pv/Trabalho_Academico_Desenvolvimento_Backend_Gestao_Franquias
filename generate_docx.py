import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Configuração de Margens
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Estilos de Cores
    NAVY = RGBColor(14, 43, 92)     # #0E2B5C UNINTER Navy
    GOLD = RGBColor(212, 138, 14)   # #D48A0E UNINTER Gold/Orange
    DARK = RGBColor(33, 37, 41)
    GRAY = RGBColor(108, 117, 125)

    # --- CAPA ---
    title_uni = doc.add_paragraph()
    title_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = title_uni.add_run("CENTRO UNIVERSITÁRIO INTERNACIONAL UNINTER\n")
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = NAVY

    r2 = title_uni.add_run("CURSO SUPERIOR DE TECNOLOGIA EM ANÁLISE E DESENVOLVIMENTO DE SISTEMAS\nDISCIPLINA: DESENVOLVIMENTO WEB BACK-END\nANO: 2026")
    r2.font.size = Pt(11)
    r2.font.color.rgb = GRAY

    doc.add_paragraph("\n" * 4)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_main = p_title.add_run("RELATÓRIO DO TRABALHO ACADÊMICO\n")
    r_main.font.size = Pt(22)
    r_main.font.bold = True
    r_main.font.color.rgb = NAVY

    r_sub = p_title.add_run("SISTEMA DE GESTÃO DE FRANQUIAS\nAPI REST EM C# E ASP.NET CORE COM ENTITY FRAMEWORK CORE")
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = GOLD

    doc.add_paragraph("\n" * 6)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_info.paragraph_format.line_spacing = 1.3
    
    r_info = p_info.add_run(
        "Estudante: WALLACE F G SILVA\n"
        "RU: 5146520\n"
        "Professor Orientador: Prof. Rodrigo da S. do Nascimento\n"
        "Ambiente: C# 12 / .NET 8.0 Web API / SQLite\n"
        "Data de Entrega: Setembro de 2026\n"
    )
    r_info.font.size = Pt(12)
    r_info.font.bold = True
    r_info.font.color.rgb = DARK

    doc.add_page_break()

    # Função Auxiliar de Título de Seção
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title)
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = NAVY

    def add_sub_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = GOLD

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(40, 40, 40)
        doc.add_paragraph() # espaçador

    # --- SEÇÃO 1: INTRODUÇÃO E PROBLEMA ---
    add_section_header("1. Introdução e Descrição do Problema")
    add_body_p(
        "No modelo empresarial de franchising, a integridade operacional e a padronização de rotinas entre a franqueadora (matriz) "
        "e suas unidades franqueadas constituem o alicerce para a expansão sustentável da marca. Contudo, na ausência de uma plataforma "
        "digital centralizada, as franquias frequentemente operam de forma isolada, gerando dispersão de informações em planilhas "
        "eletrônicas, lentidão na comunicação e ausência de conciliação fiscal e financeira."
    )
    add_body_p(
        "Para solucionar este desafio corporativo, o presente trabalho acadêmico apresenta o desenvolvimento de uma API REST completa "
        "em C# com ASP.NET Core Web API e Entity Framework Core. A solução centraliza o cadastro das unidades e seus respectivos "
        "responsáveis, mantém o catálogo padronizado de produtos e serviços, rastreia fornecedores homologados, monitora estoques "
        "em tempo real impedindo saldos negativos, orquestra as vendas em balcão com baixa atômica, calcula automaticamente os royalties "
        "devidos por competência e fornece indicadores gerenciais detalhados para tomada de decisão estratégica pela diretoria da rede."
    )

    # --- SEÇÃO 2: ARQUITETURA DO SISTEMA ---
    add_section_header("2. Arquitetura e Organização do Projeto")
    add_body_p(
        "O sistema adota uma arquitetura em camadas orientada a objetos, fundamentada no princípio da responsabilidade única (SRP) "
        "e inversão de controle (IoC) nativa do .NET 8:"
    )

    arch_points = [
        ("Controllers (Franquias.Api/Controllers): ", "Exposição dos endpoints RESTful, validação de payload, autenticação e documentação OpenAPI/Swagger."),
        ("Services (Franquias.Api/Services): ", "Camada de domínio e regras de negócio corporativas (cálculo de royalties, validação de unidades inativas, controle transacional de vendas e bloqueio de estoque negativo)."),
        ("Repositories (Franquias.Api/Repositories): ", "Abstração de acesso a dados (Repository Pattern) utilizando LINQ otimizado e Entity Framework Core."),
        ("Data & Migrations (Franquias.Api/Data): ", "Mapeamento relacional via Fluent API no AppDbContext, histórico de Migrations e classe DbInitializer para seed completo."),
        ("DTOs (Franquias.Api/DTOs): ", "Objetos de transferência que blindam o modelo interno e evitam a exposição de dados sensíveis."),
        ("Configurations (Franquias.Api/Configurations): ", "Injeção de dependência modular, autenticação JWT Bearer e Middleware global para captura e padronização de exceções HTTP.")
    ]

    for title, desc in arch_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(title)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r2 = p.add_run(desc)
        r2.font.color.rgb = DARK

    # --- SEÇÃO 3: MODELO RELACIONAL E BANCO DE DADOS ---
    add_section_header("3. Modelagem Relacional do Banco de Dados")
    add_body_p(
        "O modelo relacional foi estruturado para refletir a totalidade das operações da franquia. Foi adotado o SQLite "
        "com Entity Framework Core, garantindo conformidade com transações ACID, integridade referencial, chaves primárias e índices únicos, "
        "além de permitir a portabilidade total da aplicação sem exigir a instalação de instâncias externas de banco pelo professor avaliador."
    )

    # Tabela de Entidades
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Tabela / Entidade"
    hdr_cells[1].text = "Chaves e Índices"
    hdr_cells[2].text = "Papel no Sistema"

    for c in hdr_cells:
        set_cell_background(c, "0E2B5C")
        for p in c.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    entities_data = [
        ("Franqueadoras", "PK: Id | Unique: CNPJ", "Dados da matriz da rede e percentual padrão de royalties."),
        ("Responsaveis", "PK: Id", "Dados cadastrais e contato dos franqueados responsáveis legais."),
        ("Unidades", "PK: Id | FK: FranqueadoraId, ResponsavelId | Unique: CNPJ", "Lojas físicas, endereços, situação (Ativa/Inativa) e alíquota de royalty."),
        ("Usuarios", "PK: Id | FK: UnidadeId | Unique: Email", "Operadores, Gestores e Administradores com perfis e senha com hash BCrypt."),
        ("Categorias", "PK: Id", "Categorização dos produtos e serviços padronizados."),
        ("Fornecedores", "PK: Id | Unique: CNPJ", "Fornecedores homologados pela rede."),
        ("Produtos", "PK: Id | FK: CategoriaId, FornecedorId | Unique: CodigoSKU", "Catálogo padrão de mercadorias e serviços com preço base e estoque mínimo."),
        ("Estoques", "PK: Id | FK: UnidadeId, ProdutoId | Unique: (UnidadeId, ProdutoId)", "Saldo físico disponível por produto em cada unidade franqueada."),
        ("MovimentacoesEstoque", "PK: Id | FK: UnidadeId, ProdutoId, UsuarioId", "Histórico de auditoria de todas as entradas, saídas e ajustes de saldo."),
        ("Vendas", "PK: Id | FK: UnidadeId, UsuarioId | Unique: CodigoVenda", "Registro transacional do cupom de venda com total consolidado."),
        ("ItensVenda", "PK: Id | FK: VendaId, ProdutoId", "Itens comercializados com snapshot de preço unitário e subtotal."),
        ("Royalties", "PK: Id | FK: UnidadeId | Unique: (UnidadeId, Mes, Ano)", "Apuração financeira de royalties sobre faturamento bruto do período."),
        ("Chamados", "PK: Id | FK: UnidadeId, UsuarioAberturaId", "Solicitações de suporte operacional e técnico entre franquia e matriz.")
    ]

    for ent, keys, role in entities_data:
        row_cells = table.add_row().cells
        row_cells[0].text = ent
        row_cells[1].text = keys
        row_cells[2].text = role
        for i, c in enumerate(row_cells):
            set_cell_background(c, "F8F9FA" if len(table.rows) % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(9)

    doc.add_paragraph() # separador

    # --- SEÇÃO 4: REGRAS DE NEGÓCIO ---
    add_section_header("4. Principais Regras de Negócio Implementadas")

    rules = [
        ("1. Unicidade de CNPJ de Unidades: ", "O sistema valida tanto em memória quanto no banco relacional (Index Unique) a unicidade de CNPJs de unidades e fornecedores, impedindo cadastros duplicados com retorno HTTP 400 Bad Request."),
        ("2. Bloqueio de Operações em Unidades Inativas: ", "Unidades franqueadas que estejam com situação Inativa têm seu acesso bloqueado para emissão de vendas (retorno HTTP 400 descritivo)."),
        ("3. Impedimento Estrito de Saldo Negativo de Estoque: ", "Nenhuma movimentação de saída manual ou conclusão de venda pode deixar o saldo de estoque inferior a zero. Caso a quantidade demandada supere a disponível, a operação é interrompida."),
        ("4. Transação Atômica de Venda com Baixa de Estoque: ", "Ao registrar uma venda, o sistema abre uma transação via EF Core que persiste o cupom da venda, registra os itens e debita simultaneamente os saldos de estoque correspondentes."),
        ("5. Cálculo e Apuração Dinâmica de Royalties: ", "A rotina de apuração consolida todas as vendas com status Concluída dentro do mês e ano informados e aplica a alíquota contratual da unidade franqueada, registrando a data de vencimento para o dia 10 do mês subsequente."),
        ("6. Autenticação e Perfis de Acesso (RBAC): ", "Usuários autenticam-se com token JWT com validade de 8 horas. Endpoints críticos são protegidos por perfis: AdminFranqueadora (acesso amplo), GestorUnidade (gestão da sua unidade) e Operador (rotinas de PDV)."),
        ("7. Tratamento Centralizado de Exceções: ", "O ExceptionHandlingMiddleware captura erros não tratados e formata retornos padronizados com os status HTTP adequados (400, 401, 403, 404 e 500).")
    ]

    for r_title, r_desc in rules:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        r_t = p.add_run(r_title)
        r_t.font.bold = True
        r_t.font.color.rgb = NAVY
        r_d = p.add_run(r_desc)
        r_d.font.color.rgb = DARK

    # --- SEÇÃO 5: EVIDÊNCIAS DE EXECUÇÃO E TESTES ---
    add_section_header("5. Evidências de Execução dos Principais Endpoints")
    add_body_p(
        "A seguir são apresentadas as evidências reais de execução dos testes automatizados efetuados na API, "
        "demonstrando o comportamento correto de cada funcionalidade e regra de negócio exigida:"
    )

    evidences = [
        ("Evidência 1: Autenticação com Emissão de Token JWT",
         "POST /api/auth/login\n"
         "Request: {\"email\": \"admin@franquias.com.br\", \"senha\": \"Admin@123\"}\n\n"
         "Response HTTP 200 OK:\n"
         "{\n"
         "  \"token\": \"eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...\",\n"
         "  \"expiracao\": \"2026-09-14T14:47:18Z\",\n"
         "  \"usuario\": {\n"
         "    \"id\": 1,\n"
         "    \"nome\": \"Wallace Silva (Administrador)\",\n"
         "    \"email\": \"admin@franquias.com.br\",\n"
         "    \"perfil\": \"AdminFranqueadora\",\n"
         "    \"ativo\": true\n"
         "  }\n"
         "}"),

        ("Evidência 2: Bloqueio de Unidade com CNPJ Duplicado",
         "POST /api/unidades\n"
         "Payload com CNPJ já existente: 23.456.789/0001-01\n\n"
         "Response HTTP 400 Bad Request:\n"
         "{\n"
         "  \"sucesso\": false,\n"
         "  \"codigoStatus\": 400,\n"
         "  \"mensagem\": \"Já existe uma unidade franqueada cadastrada com o CNPJ '23.456.789/0001-01'.\",\n"
         "  \"dataHora\": \"2026-09-14T06:47:20Z\"\n"
         "}"),

        ("Evidência 3: Bloqueio de Venda para Unidade Inativa",
         "POST /api/vendas\n"
         "Tentativa de registrar venda para a Unidade Barra RJ (Id 3, Situacao = Inativa)\n\n"
         "Response HTTP 400 Bad Request:\n"
         "{\n"
         "  \"sucesso\": false,\n"
         "  \"codigoStatus\": 400,\n"
         "  \"mensagem\": \"Não é permitido registrar vendas para a unidade 'Café Gourmet - Unidade Barra RJ' pois ela se encontra INATIVA.\",\n"
         "  \"dataHora\": \"2026-09-14T06:47:21Z\"\n"
         "}"),

        ("Evidência 4: Bloqueio de Movimentação que Resulte em Saldo Negativo",
         "POST /api/estoques/movimentar\n"
         "Tentativa de SaidaVenda de 99.999 unidades do Café Espresso Blend\n\n"
         "Response HTTP 400 Bad Request:\n"
         "{\n"
         "  \"sucesso\": false,\n"
         "  \"codigoStatus\": 400,\n"
         "  \"mensagem\": \"Operação não permitida. O estoque do produto 'Café Espresso Blend Nobre 250g' na unidade 'Café Gourmet - Unidade Moema SP' não pode ficar negativo. Saldo disponível: 50, Quantidade solicitada: 99999.\",\n"
         "  \"dataHora\": \"2026-09-14T06:47:21Z\"\n"
         "}"),

        ("Evidência 5: Registro de Venda Válida com Atualização de Estoque",
         "POST /api/vendas\n"
         "Payload com 2 cafés blend (R$ 38,50 cada) e 1 croissant (R$ 15,00)\n\n"
         "Response HTTP 201 Created:\n"
         "{\n"
         "  \"id\": 4,\n"
         "  \"codigoVenda\": \"VND-20260914-4BD8EE\",\n"
         "  \"unidadeNome\": \"Café Gourmet - Unidade Moema SP\",\n"
         "  \"usuarioNome\": \"Lucas Santos (Operador SP)\",\n"
         "  \"valorTotal\": 92.00,\n"
         "  \"status\": \"Concluida\",\n"
         "  \"itens\": [\n"
         "    { \"produtoNome\": \"Café Espresso Blend Nobre 250g\", \"quantidade\": 2, \"subtotal\": 77.00 },\n"
         "    { \"produtoNome\": \"Croissant Francês Folhado Tradicional\", \"quantidade\": 1, \"subtotal\": 15.00 }\n"
         "  ]\n"
         "}"),

        ("Evidência 6: Apuração de Royalties sobre o Faturamento da Unidade",
         "POST /api/royalties/apurar\n"
         "Request: { \"unidadeId\": 1, \"mesReferencia\": 9, \"anoReferencia\": 2026 }\n\n"
         "Response HTTP 200 OK:\n"
         "{\n"
         "  \"id\": 3,\n"
         "  \"unidadeNome\": \"Café Gourmet - Unidade Moema SP\",\n"
         "  \"mesReferencia\": 9,\n"
         "  \"anoReferencia\": 2026,\n"
         "  \"faturamentoBase\": 624.00,\n"
         "  \"percentualCobrado\": 5.0,\n"
         "  \"valorRoyalty\": 31.20,\n"
         "  \"status\": \"Pendente\",\n"
         "  \"dataVencimento\": \"2026-10-10T00:00:00Z\"\n"
         "}"),

        ("Evidência 7: Indicadores Gerenciais e Ranking de Unidades por Faturamento",
         "GET /api/relatorios/ranking-unidades\n\n"
         "Response HTTP 200 OK:\n"
         "[\n"
         "  {\n"
         "    \"posicao\": 1,\n"
         "    \"codigoUnidade\": \"FRANQ-SP01\",\n"
         "    \"nomeUnidade\": \"Café Gourmet - Unidade Moema SP\",\n"
         "    \"cidade\": \"São Paulo\",\n"
         "    \"totalVendas\": 4,\n"
         "    \"totalFaturado\": 624.00\n"
         "  },\n"
         "  {\n"
         "    \"posicao\": 2,\n"
         "    \"codigoUnidade\": \"FRANQ-PR01\",\n"
         "    \"nomeUnidade\": \"Café Gourmet - Unidade Batel Curitiba\",\n"
         "    \"cidade\": \"Curitiba\",\n"
         "    \"totalVendas\": 1,\n"
         "    \"totalFaturado\": 115.50\n"
         "  }\n"
         "]")
    ]

    for ev_title, ev_content in evidences:
        add_sub_header(ev_title)
        add_code_block(ev_content)

    # --- SEÇÃO 6: ANÁLISE CRÍTICA ---
    add_section_header("6. Análise Crítica do Desenvolvimento")

    add_sub_header("6.1 Funcionalidades Implementadas e Destaques Técnicos")
    add_body_p(
        "Foram desenvolvidos com sucesso todos os 9 módulos funcionais exigidos na especificação: Autenticação JWT, "
        "Cadastro de Franqueadora e Unidades com validação de CNPJ, Catálogo de Produtos e Categorias, Controle de Fornecedores, "
        "Saldo e Movimentações de Estoque, Emissão de Vendas com integridade atômica, Apuração e Liquidação de Royalties, "
        "Abertura e Resolução de Chamados de Suporte, e Indicadores Analíticos com LINQ."
    )

    add_sub_header("6.2 Decisões Técnicas e Desafios Superados")
    add_body_p(
        "A decisão de utilizar o SQLite como banco relacional provou-se extremamente acertada para o contexto acadêmico, pois garante "
        "uma avaliação limpa e imediata pelo professor sem atritos de dependência de contêineres Docker ou serviços locais de terceiros. "
        "Durante a implementação dos relatórios gerenciais, superou-se a limitação do provider SQLite com o operador de agregação Sum em "
        "campos de ponto fixo (decimal), aplicando-se projeções otimizadas de colunas seguidas de cálculo em memória com LINQ to Objects."
    )

    add_sub_header("6.3 Limitações e Propostas de Evolução")
    add_body_p(
        "Como melhorias planejadas para versões futuras da plataforma corporativa, destacam-se: implementação de mensageria assíncrona "
        "(RabbitMQ ou Azure Service Bus) para notificação imediata de estoques críticos e chamados abertos aos gestores da franqueadora, "
        "e a integração de um gateway de pagamentos para geração automática de QR Codes PIX e boletos bancários nos lançamentos de royalties."
    )

    # --- SEÇÃO 7: CONCLUSÃO ---
    add_section_header("7. Conclusão e Dados de Acesso para Avaliação")
    add_body_p(
        "O projeto cumpre com rigor os requisitos acadêmicos da UNINTER, apresentando uma API REST corporativa moderna, fortemente tipada, "
        "segura, testada e completamente documentada via Swagger/OpenAPI."
    )

    add_sub_header("Resumo dos Usuários Pré-Cadastrados para Teste:")
    table_users = doc.add_table(rows=1, cols=4)
    table_users.alignment = WD_TABLE_ALIGNMENT.CENTER
    uhdr = table_users.rows[0].cells
    uhdr[0].text = "Perfil"
    uhdr[1].text = "E-mail"
    uhdr[2].text = "Senha"
    uhdr[3].text = "Escopo"

    for c in uhdr:
        set_cell_background(c, "0E2B5C")
        for p in c.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    u_data = [
        ("AdminFranqueadora", "admin@franquias.com.br", "Admin@123", "Acesso irrestrito a toda a rede"),
        ("GestorUnidade", "gestor.sp@franquias.com.br", "Gestor@123", "Unidade Moema SP"),
        ("Operador", "operador.sp@franquias.com.br", "Operador@123", "PDV Unidade Moema SP"),
        ("GestorUnidade", "gestor.pr@franquias.com.br", "Gestor@123", "Unidade Batel Curitiba"),
        ("Operador", "operador.pr@franquias.com.br", "Operador@123", "PDV Unidade Curitiba")
    ]

    for p_val, e_val, s_val, esc_val in u_data:
        r_cells = table_users.add_row().cells
        r_cells[0].text = p_val
        r_cells[1].text = e_val
        r_cells[2].text = s_val
        r_cells[3].text = esc_val
        for c in r_cells:
            set_cell_background(c, "F8F9FA" if len(table_users.rows) % 2 == 0 else "FFFFFF")
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(9.5)

    doc.add_paragraph()
    p_final = doc.add_paragraph()
    p_final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_final.add_run("Trabalho Acadêmico — WALLACE F G SILVA (RU: 5146520) — UNINTER 2026")
    r_f.font.bold = True
    r_f.font.size = Pt(10)
    r_f.font.color.rgb = GRAY

    output_path = os.path.join(os.getcwd(), "docs", "Relatorio_Academico_Gestao_Franquias.docx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Documento DOCX salvo com sucesso em: {output_path}")

if __name__ == "__main__":
    create_report()
